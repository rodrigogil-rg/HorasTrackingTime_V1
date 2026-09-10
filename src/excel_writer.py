import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.models import TimeEntry
from src.transformer import group_by_client

SPANISH_MONTHS = {
    1: "Enero",
    2: "Febrero",
    3: "Marzo",
    4: "Abril",
    5: "Mayo",
    6: "Junio",
    7: "Julio",
    8: "Agosto",
    9: "Septiembre",
    10: "Octubre",
    11: "Noviembre",
    12: "Diciembre",
}


def sanitize_sheet_name(name: str, existing_names: set[str]) -> str:
    invalid = r"[\\/*?:\[\]]"
    clean = re.sub(invalid, "_", name).strip()
    if not clean:
        clean = "Cliente"

    clean = clean[:31]

    candidate = clean
    counter = 2
    while candidate in existing_names:
        suffix = f"_{counter}"
        max_len = 31 - len(suffix)
        candidate = clean[:max_len] + suffix
        counter += 1

    return candidate


def sanitize_filename(name: str) -> str:
    invalid = r'[\\/*?:\[\]<>|"]'
    clean = re.sub(invalid, "_", name).strip()
    return re.sub(r"\s+", "_", clean)


def write_excel(
    entries: list[TimeEntry], mes: int, anio: int, usuario: str, output_dir: Path
) -> Path:
    wb = openpyxl.Workbook()
    default_sheet = wb.active

    grouped = group_by_client(entries)
    mes_nombre = SPANISH_MONTHS.get(mes, str(mes))

    # Color Palette from Golden Master Reference
    COLOR_PRIMARY_DARK = "17365D"  # Dark Blue for main title
    COLOR_HEADER_BG = "1F4E78"  # Medium Blue for table headers
    COLOR_TOTAL_BG = "F4B183"  # Orange/Peach for total rows
    COLOR_ROW_ALT = "F4F8FC"  # Very light blue tint for data rows
    COLOR_SUBTITLE_BG = "D9EAF7"  # Light blue for subtitle in client sheets
    COLOR_WHITE = "FFFFFF"

    # Fonts
    font_title_summary = Font(name="Calibri", size=16, bold=True, color=COLOR_WHITE)
    font_title_client = Font(name="Calibri", size=15, bold=True, color=COLOR_WHITE)
    font_subtitle = Font(name="Calibri", size=11, bold=True, color=COLOR_PRIMARY_DARK)
    font_header = Font(name="Calibri", size=11, bold=True, color=COLOR_WHITE)
    font_data = Font(name="Calibri", size=11, bold=False, color="000000")
    font_total = Font(name="Calibri", size=11, bold=True, color=COLOR_PRIMARY_DARK)

    # Fills
    fill_title = PatternFill(
        start_color=COLOR_PRIMARY_DARK, end_color=COLOR_PRIMARY_DARK, fill_type="solid"
    )
    fill_header = PatternFill(
        start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid"
    )
    fill_total = PatternFill(
        start_color=COLOR_TOTAL_BG, end_color=COLOR_TOTAL_BG, fill_type="solid"
    )
    fill_row_alt = PatternFill(
        start_color=COLOR_ROW_ALT, end_color=COLOR_ROW_ALT, fill_type="solid"
    )
    fill_subtitle = PatternFill(
        start_color=COLOR_SUBTITLE_BG, end_color=COLOR_SUBTITLE_BG, fill_type="solid"
    )

    # Borders
    thin_border_side = Side(border_style="thin", color="D9D9D9")
    thin_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side,
    )

    total_border = Border(
        left=Side(border_style="thin", color="D9D9D9"),
        right=Side(border_style="thin", color="D9D9D9"),
        top=Side(border_style="thin", color="D9D9D9"),
        bottom=Side(border_style="double", color="000000"),
    )

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # 1. Resumen general
    ws_summary = wb.create_sheet(title="Resumen general")

    # Title A1:E1
    ws_summary.merge_cells("A1:E1")
    title_cell = ws_summary["A1"]
    title_cell.value = f"Resumen general de horas - {usuario} - {mes_nombre} {anio}"
    title_cell.font = font_title_summary
    title_cell.fill = fill_title
    title_cell.alignment = align_center
    ws_summary.row_dimensions[1].height = 30

    # Row 3: Headers
    headers_summary = [
        "Cliente",
        "Horas totales",
        "Registros",
        "Días trabajados",
        "% del total",
    ]
    ws_summary.row_dimensions[3].height = 20
    for col_idx, h in enumerate(headers_summary, start=1):
        cell = ws_summary.cell(row=3, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = thin_border

    current_row = 4
    total_row_idx = 3 + len(grouped) + 1

    for client in sorted(grouped.keys()):
        ws_summary.row_dimensions[current_row].height = 18

        # A: client
        cell_a = ws_summary.cell(row=current_row, column=1, value=client)
        cell_a.font = font_data
        cell_a.fill = fill_row_alt
        cell_a.alignment = align_left
        cell_a.border = thin_border

        # B: Horas totales
        cell_b = ws_summary.cell(row=current_row, column=2, value=f"='{client}'!K2")
        cell_b.font = font_data
        cell_b.fill = fill_row_alt
        cell_b.alignment = align_right
        cell_b.number_format = "0.00"
        cell_b.border = thin_border

        # C: Registros
        client_entries = grouped[client]
        last_client_row = 3 + len(client_entries)
        cell_c = ws_summary.cell(
            row=current_row, column=3, value=f"=ROWS('{client}'!A4:A{last_client_row})"
        )
        cell_c.font = font_data
        cell_c.fill = fill_row_alt
        cell_c.alignment = align_right
        cell_c.number_format = "#,##0"
        cell_c.border = thin_border

        # D: Días trabajados
        cell_d = ws_summary.cell(
            row=current_row,
            column=4,
            value=f"=COUNTA(_xlfn.UNIQUE('{client}'!G4:G{last_client_row}))",
        )
        cell_d.font = font_data
        cell_d.fill = fill_row_alt
        cell_d.alignment = align_right
        cell_d.number_format = "#,##0"
        cell_d.border = thin_border

        # E: % del total
        cell_e = ws_summary.cell(
            row=current_row, column=5, value=f"=B{current_row}/$B${total_row_idx}"
        )
        cell_e.font = font_data
        cell_e.fill = fill_row_alt
        cell_e.alignment = align_right
        cell_e.number_format = "0.0%"
        cell_e.border = thin_border

        current_row += 1

    # Total General Row
    total_row = current_row
    ws_summary.row_dimensions[total_row].height = 20

    cell_tot_a = ws_summary.cell(row=total_row, column=1, value="TOTAL GENERAL")
    cell_tot_a.font = font_total
    cell_tot_a.fill = fill_total
    cell_tot_a.alignment = align_left
    cell_tot_a.border = total_border

    cell_tot_b = ws_summary.cell(row=total_row, column=2, value=f"=SUM(B4:B{total_row - 1})")
    cell_tot_b.font = font_total
    cell_tot_b.fill = fill_total
    cell_tot_b.alignment = align_right
    cell_tot_b.number_format = "0.00"
    cell_tot_b.border = total_border

    cell_tot_c = ws_summary.cell(row=total_row, column=3, value=f"=SUM(C4:C{total_row - 1})")
    cell_tot_c.font = font_total
    cell_tot_c.fill = fill_total
    cell_tot_c.alignment = align_right
    cell_tot_c.number_format = "#,##0"
    cell_tot_c.border = total_border

    cell_tot_d = ws_summary.cell(row=total_row, column=4, value="")
    cell_tot_d.font = font_total
    cell_tot_d.fill = fill_total
    cell_tot_d.border = total_border

    cell_tot_e = ws_summary.cell(row=total_row, column=5, value=f"=SUM(E4:E{total_row - 1})")
    cell_tot_e.font = font_total
    cell_tot_e.fill = fill_total
    cell_tot_e.alignment = align_right
    cell_tot_e.number_format = "0.0%"
    cell_tot_e.border = total_border

    for col in ws_summary.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_summary.column_dimensions[col_letter].width = max(max_len + 5, 18)

    # 2. Client Sheets
    existing_sheet_names = {"Resumen general"}
    for client, client_entries in grouped.items():
        sheet_name = sanitize_sheet_name(client, existing_sheet_names)
        existing_sheet_names.add(sheet_name)

        ws = wb.create_sheet(title=sheet_name)

        # Title A1:K1
        ws.merge_cells("A1:K1")
        t_cell = ws["A1"]
        t_cell.value = f"Horas de {usuario} - {client} - {mes_nombre} {anio}"
        t_cell.font = font_title_client
        t_cell.fill = fill_title
        t_cell.alignment = align_center
        ws.row_dimensions[1].height = 30

        # Row 2: A2:H2 merged "Detalle de horas informadas por cliente", I2:J2 merged "Total de horas", K2 formula
        ws.merge_cells("A2:H2")
        sub_cell = ws["A2"]
        sub_cell.value = "Detalle de horas informadas por cliente"
        sub_cell.font = font_subtitle
        sub_cell.fill = fill_subtitle
        sub_cell.alignment = align_left
        for col in range(1, 9):
            ws.cell(row=2, column=col).fill = fill_subtitle
            ws.cell(row=2, column=col).border = thin_border

        ws.merge_cells("I2:J2")
        tot_label_cell = ws["I2"]
        tot_label_cell.value = "Total de horas"
        tot_label_cell.font = font_subtitle
        tot_label_cell.fill = fill_subtitle
        tot_label_cell.alignment = align_right
        for col in range(9, 11):
            ws.cell(row=2, column=col).fill = fill_subtitle
            ws.cell(row=2, column=col).border = thin_border

        last_data_row = 3 + len(client_entries)
        k2_cell = ws["K2"]
        k2_cell.value = f"=SUM(I4:I{last_data_row})"
        k2_cell.font = font_total
        k2_cell.fill = fill_total
        k2_cell.alignment = align_right
        k2_cell.number_format = "0.00"
        k2_cell.border = total_border
        ws.row_dimensions[2].height = 20

        # Row 3: Headers
        headers_client = [
            "Usuario",
            "Proyecto",
            "Tarea",
            "Cliente",
            "Servicio",
            "Notas",
            "Fecha",
            "Duración",
            "Horas",
            "Desde",
            "Hasta",
        ]
        ws.row_dimensions[3].height = 20
        for col_idx, h in enumerate(headers_client, start=1):
            cell = ws.cell(row=3, column=col_idx, value=h)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = thin_border

        # Data rows
        for idx, entry in enumerate(client_entries, start=4):
            ws.row_dimensions[idx].height = 18
            # A: Usuario
            c_a = ws.cell(row=idx, column=1, value=entry.usuario)
            c_a.font = font_data
            c_a.alignment = align_left
            c_a.border = thin_border

            # B: Proyecto
            c_b = ws.cell(row=idx, column=2, value=entry.proyecto)
            c_b.font = font_data
            c_b.alignment = align_left
            c_b.border = thin_border

            # C: Tarea
            c_c = ws.cell(row=idx, column=3, value=entry.tarea)
            c_c.font = font_data
            c_c.alignment = align_left
            c_c.border = thin_border

            # D: Cliente
            c_d = ws.cell(row=idx, column=4, value=entry.cliente)
            c_d.font = font_data
            c_d.alignment = align_left
            c_d.border = thin_border

            # E: Servicio
            c_e = ws.cell(row=idx, column=5, value=entry.servicio or "")
            c_e.font = font_data
            c_e.alignment = align_left
            c_e.border = thin_border

            # F: Notas
            c_f = ws.cell(row=idx, column=6, value=entry.notas)
            c_f.font = font_data
            c_f.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            c_f.border = thin_border

            # G: Fecha
            c_g = ws.cell(row=idx, column=7, value=entry.fecha.strftime("%Y-%m-%d"))
            c_g.font = font_data
            c_g.alignment = align_center
            c_g.border = thin_border

            # H: Duración
            c_h = ws.cell(row=idx, column=8, value=entry.duracion_str)
            c_h.font = font_data
            c_h.alignment = align_center
            c_h.border = thin_border

            # I: Horas
            c_i = ws.cell(row=idx, column=9, value=entry.horas)
            c_i.font = font_data
            c_i.alignment = align_right
            c_i.number_format = "0.00"
            c_i.border = thin_border

            # J: Desde
            c_j = ws.cell(row=idx, column=10, value=entry.hora_inicio.strftime("%H:%M:%S"))
            c_j.font = font_data
            c_j.alignment = align_center
            c_j.border = thin_border

            # K: Hasta
            c_k = ws.cell(row=idx, column=11, value=entry.hora_fin.strftime("%H:%M:%S"))
            c_k.font = font_data
            c_k.alignment = align_center
            c_k.border = thin_border

        ws.freeze_panes = "A4"
        ws.auto_filter.ref = f"A3:K{last_data_row}"

        col_widths = {
            "A": 15,
            "B": 16,
            "C": 52,
            "D": 15,
            "E": 15,
            "F": 48,
            "G": 13,
            "H": 13,
            "I": 13,
            "J": 13,
            "K": 13,
        }
        for col_letter, width in col_widths.items():
            ws.column_dimensions[col_letter].width = width

    if default_sheet in wb.worksheets and len(wb.worksheets) > 1:
        wb.remove(default_sheet)

    output_dir.mkdir(parents=True, exist_ok=True)
    filename = f"Planilla_Horas_{sanitize_filename(usuario)}_{mes_nombre}_{anio}.xlsx"
    output_path = output_dir / filename

    wb.save(output_path)
    return output_path
