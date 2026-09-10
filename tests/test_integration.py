from pathlib import Path

import openpyxl

from src.excel_writer import write_excel
from src.parser import parse_csv_file
from src.transformer import filter_by_month, group_by_client


def test_parser_and_validators():
    csv_path = Path("Ejemplos/Planilla de tiempo, ago 1,2026 - ago 31,2026 - 80643.csv")
    entries = parse_csv_file(csv_path)
    assert len(entries) == 23
    assert entries[0].cliente == "Cirion"
    assert entries[0].horas == 1.0
    assert entries[0].duracion_timedelta.total_seconds() == 3600


def test_transformer_monthly_filter():
    csv_path = Path("Ejemplos/Planilla de tiempo, ago 1,2026 - ago 31,2026 - 80643.csv")
    entries = parse_csv_file(csv_path)
    filtered = filter_by_month(entries, 8, 2026)
    assert len(filtered) == 23

    filtered_empty = filter_by_month(entries, 7, 2026)
    assert len(filtered_empty) == 0


def test_grouping_and_sorting():
    csv_path = Path("Ejemplos/Planilla de tiempo, ago 1,2026 - ago 31,2026 - 80643.csv")
    entries = parse_csv_file(csv_path)
    filtered = filter_by_month(entries, 8, 2026)
    grouped = group_by_client(filtered)
    assert "Cirion" in grouped
    assert "Mirgor" in grouped
    assert len(grouped["Cirion"]) == 22
    assert len(grouped["Mirgor"]) == 1


def test_excel_writer_generation(tmp_path):
    csv_path = Path("Ejemplos/Planilla de tiempo, ago 1,2026 - ago 31,2026 - 80643.csv")
    entries = parse_csv_file(csv_path)
    filtered = filter_by_month(entries, 8, 2026)
    out_path = write_excel(filtered, 8, 2026, "Rodrigo Gil", tmp_path)
    assert out_path.exists()

    wb = openpyxl.load_workbook(out_path, data_only=False)
    assert "Resumen general" in wb.sheetnames
    assert "Cirion" in wb.sheetnames
    assert "Mirgor" in wb.sheetnames
