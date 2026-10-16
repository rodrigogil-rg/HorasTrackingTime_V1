from datetime import datetime

import openpyxl

from src.excel_writer import sanitize_filename, sanitize_sheet_name, write_excel
from src.models import TimeEntry


def test_sanitize_sheet_name():
    """Prueba la sanitización de nombres de hojas de Excel: caracteres inválidos, colisiones (Cliente_A_2) y límite de 31 caracteres."""
    existing = {"Resumen general"}
    assert sanitize_sheet_name("Cliente A", existing) == "Cliente A"
    existing.add("Cliente A")
    assert sanitize_sheet_name("Cliente A", existing) == "Cliente A_2"
    assert sanitize_sheet_name("Cliente/Con:Barras", existing) == "Cliente_Con_Barras"

    # Max 31 chars test
    long_name = "A" * 40
    sanitized = sanitize_sheet_name(long_name, set())
    assert len(sanitized) <= 31


def test_sanitize_filename():
    """Prueba la sanitización de nombres de archivo de salida: reemplaza espacios y caracteres inválidos por guiones bajos."""
    assert sanitize_filename("Rodrigo Gil") == "Rodrigo_Gil"
    assert sanitize_filename("Test / Name*?:") == "Test___Name___"


def test_excel_writer_total_dias_trabajados_unicos(tmp_path):
    """Prueba que la fila TOTAL GENERAL calcule días trabajados únicos en todos los registros.

    Si dos clientes tienen registros en la misma fecha, ese día debe contarse una sola vez
    en el total general (no sumarse dos veces).
    """
    entries = [
        TimeEntry(
            servicio=None,
            cliente="Cliente A",
            proyecto="Proyecto A",
            usuario="Usuario Test",
            lista_tareas=None,
            tarea="Tarea 1",
            fecha_entrega=None,
            horas_estimadas=None,
            facturable=True,
            facturado=False,
            precio_hora=None,
            archivado=False,
            fecha_inicio=datetime(2026, 8, 3, 9, 0, 0),  # noqa: DTZ001
            fecha_fin=datetime(2026, 8, 3, 10, 0, 0),  # noqa: DTZ001
            zona_horaria="GMT-03:00",
            duracion_str="1:00:00",
            horas=1.0,
            notas="Registro 1",
            total=None,
            costo_hora=None,
            costo=None,
            moneda="USD",
        ),
        TimeEntry(
            servicio=None,
            cliente="Cliente A",
            proyecto="Proyecto A",
            usuario="Usuario Test",
            lista_tareas=None,
            tarea="Tarea 2",
            fecha_entrega=None,
            horas_estimadas=None,
            facturable=True,
            facturado=False,
            precio_hora=None,
            archivado=False,
            fecha_inicio=datetime(2026, 8, 3, 11, 0, 0),  # noqa: DTZ001
            fecha_fin=datetime(2026, 8, 3, 12, 0, 0),  # noqa: DTZ001
            zona_horaria="GMT-03:00",
            duracion_str="1:00:00",
            horas=1.0,
            notas="Registro 2",
            total=None,
            costo_hora=None,
            costo=None,
            moneda="USD",
        ),
        TimeEntry(
            servicio=None,
            cliente="Cliente B",
            proyecto="Proyecto B",
            usuario="Usuario Test",
            lista_tareas=None,
            tarea="Tarea 3",
            fecha_entrega=None,
            horas_estimadas=None,
            facturable=True,
            facturado=False,
            precio_hora=None,
            archivado=False,
            fecha_inicio=datetime(2026, 8, 3, 13, 0, 0),  # noqa: DTZ001
            fecha_fin=datetime(2026, 8, 3, 14, 0, 0),  # noqa: DTZ001
            zona_horaria="GMT-03:00",
            duracion_str="1:00:00",
            horas=1.0,
            notas="Registro 3",
            total=None,
            costo_hora=None,
            costo=None,
            moneda="USD",
        ),
        TimeEntry(
            servicio=None,
            cliente="Cliente B",
            proyecto="Proyecto B",
            usuario="Usuario Test",
            lista_tareas=None,
            tarea="Tarea 4",
            fecha_entrega=None,
            horas_estimadas=None,
            facturable=True,
            facturado=False,
            precio_hora=None,
            archivado=False,
            fecha_inicio=datetime(2026, 8, 4, 9, 0, 0),  # noqa: DTZ001
            fecha_fin=datetime(2026, 8, 4, 11, 0, 0),  # noqa: DTZ001
            zona_horaria="GMT-03:00",
            duracion_str="2:00:00",
            horas=2.0,
            notas="Registro 4",
            total=None,
            costo_hora=None,
            costo=None,
            moneda="USD",
        ),
    ]

    out_path = write_excel(entries, 8, 2026, "Usuario Test", tmp_path)
    assert out_path.exists()

    wb = openpyxl.load_workbook(out_path, data_only=False)
    assert "Resumen general" in wb.sheetnames

    ws = wb["Resumen general"]

    # Verificar días por cliente
    # Cliente A tiene 1 día único (03/08/2026) - dos registros mismo día
    # Cliente B tiene 2 días únicos (03/08/2026 y 04/08/2026)
    assert ws.cell(row=4, column=4).value == 1  # Cliente A: 1 día único
    assert ws.cell(row=5, column=4).value == 2  # Cliente B: 2 días únicos

    # Fila 6 es TOTAL GENERAL
    # Días únicos totales: 03/08/2026 y 04/08/2026 = 2 días únicos
    # NO 3 (que sería la suma de 1 + 2)
    assert ws.cell(row=6, column=4).value == 2  # Total: 2 días únicos, no 3
